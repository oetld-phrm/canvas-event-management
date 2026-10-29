"""Update existing Canvas calendar events from an edited export (Python 3.10+).

Install: python -m pip install requests openpyxl tzdata
Preview: python -m canvas_events update
Apply:   python -m canvas_events update --apply

The default input is the companion exporter's output CSV. The entire import is
planned before writes begin. Blank cells preserve values; __CLEAR__ empties a
supported text field. Assignment rows are automatically skipped.
"""
import argparse
import csv
import getpass
import json
import os
import re
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
import requests
from .common import canvas_request, validate_base_url, connect, resolve_input_path

FIELDS = {'title', 'description', 'start_at', 'end_at', 'location_name',
          'location_address', 'context_code', 'all_day', 'blackout_date',
          'time_zone_edited', 'rrule'}
# Canvas accepts these editable fields; export metadata is never sent back.
BOOLS = {'all_day', 'blackout_date'}
DEFAULT_CSV = 'export-events.csv'
CLEARABLE_FIELDS = {'description', 'location_name', 'location_address'}
SHARED_SECTION_FIELDS = {'title', 'description', 'location_name', 'location_address'}
IGNORED_COLUMNS = {'course_id', 'course_name', 'source', 'created_at',
                   'updated_at', 'workflow_state'}
# item_id is the exporter's name for the Canvas calendar event ID.
HEADER_ALIASES = {'item_id': 'event_id', 'new_description': 'description'}

def normalize_value(key, value):
    """Normalize both spreadsheet and Canvas values before comparing them.

    Booleans become Python bools; offset-aware timestamps become UTC ISO text.
    UTC conversion prevents equivalent times from looking like changes. Text is
    preserved verbatim here; only None is treated as an empty string.
    """
    if key in BOOLS:
        text = str(value).lower() if value is not None else 'false'
        if text not in {'true', 'false', '1', '0', 'yes', 'no'}:
            raise ValueError(f'{key}: use true or false')
        return text in {'true', '1', 'yes'}
    if key in {'start_at', 'end_at'} and value:
        dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if dt.tzinfo is None:
            raise ValueError(f'{key}: include a UTC offset, e.g. 2026-11-02T09:00:00-08:00')
        return dt.astimezone(timezone.utc).isoformat()
    return '' if value is None else str(value)

def load_spreadsheet(path, sheet=None):
    """Read cells without interpreting them; XLSX support is loaded only if used."""
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f))
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = list((book[sheet] if sheet else book.active).iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ValueError('Use .csv or .xlsx')
    return rows


def is_assignment_row(row):
    """Identify assignment exports without consulting the ignored source column.

    Assignment IDs and calendar IDs can overlap. An assignment URL, nonempty
    submission type, or due_at marks a row that must not reach the event API.
    """
    return bool('/assignments/' in row.get('html_url', '')
                or row.get('assignment_type') or row.get('due_at'))


def read_rows(path, sheet=None):
    """Validate headers/IDs, ignore metadata, and retain spreadsheet row numbers.

    Empty cells are omitted from each dictionary. This distinction is intentional:
    a missing key means 'leave Canvas unchanged', not 'erase the current value'.
    """
    rows = load_spreadsheet(path, sheet)
    if not rows:
        raise ValueError('Empty spreadsheet')
    headers = [str(x).strip().lower() if x is not None else '' for x in rows[0]]
    headers = [HEADER_ALIASES.get(x, x) for x in headers]
    named = [x for x in headers if x]
    unknown = set(named) - FIELDS - {'event_id', 'sections', 'which'} - IGNORED_COLUMNS - {'assignment_type', 'html_url', 'due_at'}
    if unknown or len(named) != len(set(named)) or 'event_id' not in named:
        raise ValueError(f'Require item_id (or event_id) and unique supported headers. Unknown: {sorted(unknown)}')
    result, seen = [], set()
    for number, cells in enumerate(rows[1:], 2):
        if not any(x is not None and str(x).strip() for x in cells):
            continue
        row = {h: (str(v) if h in {'title', 'description', 'location_name', 'location_address'} else str(v).strip()) for h, v in zip(headers, cells)
               if h and h not in IGNORED_COLUMNS and v is not None and str(v).strip()}
        # Skip assignments before duplicate-ID checks: an assignment and an event
        # can legitimately have the same numeric ID in the export.
        if is_assignment_row(row):
            print(f'Row {number}: skipped assignment (calendar events only)')
            continue
        eid = row.get('event_id', '')
        if not re.fullmatch(r'[1-9]\d*', eid) or eid in seen:
            raise ValueError(f'Row {number}: invalid or duplicate event_id')
        seen.add(eid)
        result.append((number, row))
    return result

def normalize_section_schedules(items):
    """Index schedules by section code so list ordering cannot create differences.

    This also validates uniqueness, timestamps, and start/end ordering. It is used
    internally for reassignment and the optional complete-section-list format.
    """
    if not isinstance(items, list):
        raise ValueError('sections must be a JSON list')
    result = {}
    for item in items:
        code = item['context_code']
        if not re.fullmatch(r'course_section_[1-9]\d*', code) or code in result:
            raise ValueError('Invalid or duplicate section context_code')
        start, end = normalize_value('start_at', item['start_at']), normalize_value('end_at', item['end_at'])
        if not start or not end or datetime.fromisoformat(end) < datetime.fromisoformat(start):
            raise ValueError('Invalid section start/end times')
        result[code] = dict(context_code=code, start_at=start, end_at=end)
    return result

def find_changes(current, row):
    """Return only supplied values that differ from the current Canvas event."""
    changes = {}
    for key in sorted(FIELDS & row.keys()):
        value = row[key]
        # Blank cells were removed during parsing; clearing must be explicit.
        if value == '__CLEAR__':
            if key not in CLEARABLE_FIELDS:
                raise ValueError(f'Cannot clear {key} with __CLEAR__')
            value = ''
        value = normalize_value(key, value)
        if value != normalize_value(key, current.get(key)):
            changes[key] = value
    if 'context_code' in changes:
        if not re.fullmatch(r'(course|group|user|account)_[1-9]\d*', changes['context_code']):
            raise ValueError('Use context_code such as course_123')
        if current.get('parent_event_id') or current.get('child_events_count') or current.get('appointment_group_id') or 'sections' in row:
            raise ValueError('Canvas restricts calendar moves for section events and appointments')
    # Validate the resulting range, including an endpoint left unchanged.
    start = changes.get('start_at', current.get('start_at'))
    end = changes.get('end_at', current.get('end_at'))
    if start and end and datetime.fromisoformat(normalize_value('end_at', end)) < datetime.fromisoformat(normalize_value('start_at', start)):
        raise ValueError('end_at precedes start_at')
    if 'sections' in row:
        if current.get('parent_event_id') or not current.get('context_code', '').startswith('course_'):
            raise ValueError('Use the parent course event_id for sections')
        desired = normalize_section_schedules(json.loads(row['sections']))
        if not desired:
            raise ValueError('Empty section lists are unsupported')
        children = current.get('child_events') or []
        if len(children) != int(current.get('child_events_count') or 0):
            raise ValueError('Canvas did not return the complete section list')
        if desired != normalize_section_schedules(children):
            changes['sections'] = list(desired.values())
    which = row.get('which', 'one')
    if which not in {'one', 'all', 'following'}:
        raise ValueError('which must be one, all, or following')
    return changes, which


def plan_updates(session, base, current, row):
    """Route section content to its parent and schedules to section children."""
    original_code = current.get('context_code', '')
    desired_code = row.get('context_code', original_code)
    is_section = original_code.startswith('course_section_')
    if not is_section:
        if desired_code.startswith('course_section_'):
            raise ValueError('Converting a course-wide event to section scheduling requires the explicit sections column')
        changes, which = find_changes(current, row)
        return [(str(current['id']), current, changes, which)] if changes else []
    parent_id = current.get('parent_event_id')
    if not parent_id:
        raise ValueError('Section event has no parent_event_id; cannot safely update it')
    if not re.fullmatch(r'course_section_[1-9]\d*', desired_code):
        raise ValueError('A section event must retain a course_section_ID context_code')
    if row.get('which', 'one') != 'one' or 'rrule' in row or 'sections' in row:
        raise ValueError('Section rows support individual events only, without rrule or sections')
    parent = canvas_request(session, 'GET', f'{base}/api/v1/calendar_events/{parent_id}')
    # context_code identifies a section here. It cannot be sent as an ordinary
    # calendar move; section reassignment must rebuild the parent's child list.
    comparison_row = dict(row)
    comparison_row.pop('context_code', None)
    changes, which = find_changes(current, comparison_row)
    # Canvas copies these fields from the parent onto every sibling section.
    parent_changes = {key: changes.pop(key) for key in list(changes) if key in SHARED_SECTION_FIELDS}
    operations = []
    if desired_code != original_code:
        if set(changes) - {'start_at', 'end_at'}:
            raise ValueError('Change section assignment separately from all_day or other advanced fields')
        children = parent.get('child_events') or []
        if len(children) != int(parent.get('child_events_count') or 0):
            raise ValueError('Canvas did not return all sibling sections')
        # Sending child_event_data replaces the complete set. Start from Canvas's
        # full sibling list so changing one section never drops the other sections.
        schedules = normalize_section_schedules(children)
        if original_code not in schedules or not any(str(c['id']) == str(current['id']) for c in children):
            raise ValueError('Section event is not present in the parent section list')
        if desired_code in schedules:
            raise ValueError('Target section already has an event under this parent')
        schedule = schedules.pop(original_code)
        schedule.update(context_code=desired_code)
        schedule.update({k: v for k, v in changes.items() if k in {'start_at', 'end_at'}})
        schedules[desired_code] = schedule
        parent_changes['sections'] = list(schedules.values())
        changes = {}
    if parent_changes:
        operations.append((str(parent_id), parent, parent_changes, which))
    if changes:
        if current.get('workflow_state') == 'locked':
            raise ValueError('Canvas marks this section event locked; it cannot be edited directly')
        operations.append((str(current['id']), current, changes, which))
    return operations

def build_update_payload(changes, which):
    """Encode Canvas's nested form parameter names; do not send export metadata."""
    data = {'which': which}
    for key, value in changes.items():
        if key == 'sections':
            for i, section in enumerate(value):
                for field, item in section.items():
                    data[f'calendar_event[child_event_data][{i}][{field}]'] = item
        else:
            data[f'calendar_event[{key}]'] = str(value).lower() if isinstance(value, bool) else value
    return data

@dataclass
class EventUpdate:
    """One merged PUT request, with its original values and contributing CSV rows."""
    event_id: str
    before: dict
    changes: dict
    which: str
    rows: list = field(default_factory=list)


def merge_update(existing, incoming):
    """Combine edits to one target; reject conflicting changes before any writes."""
    if existing.which != incoming.which:
        raise ValueError(f'Conflicting recurrence scopes for event {existing.event_id}')
    for key, value in incoming.changes.items():
        if key != 'sections':
            if key in existing.changes and existing.changes[key] != value:
                raise ValueError(f'Conflicting {key} edits for event {existing.event_id}')
            existing.changes[key] = value
            continue
        # Each row's section list started from the same live snapshot. Merge only
        # its differences, not the unchanged sibling schedules it carries along.
        original = normalize_section_schedules(existing.before.get('child_events') or [])
        combined = normalize_section_schedules(existing.changes.get('sections', list(original.values())))
        desired = normalize_section_schedules(value)
        for code in original.keys() | desired.keys():
            old, new = original.get(code), desired.get(code)
            if old == new:
                continue
            if combined.get(code) != old and combined.get(code) != new:
                raise ValueError(f'Conflicting section edits for {code}')
            if new is None:
                combined.pop(code, None)
            else:
                combined[code] = new
        existing.changes['sections'] = list(combined.values())
    existing.rows.extend(incoming.rows)


def build_import_plan(session, base, rows):
    """Fetch live events and validate every row before permitting any PUT requests.

    Only changed shared fields participate in merging. An unchanged sibling row
    therefore cannot undo a description edit made in another section's row.
    """
    updates = {}
    unchanged = 0
    errors = []
    for number, row in rows:
        event_id = row['event_id']
        try:
            current = canvas_request(session, 'GET', f'{base}/api/v1/calendar_events/{event_id}')
            operations = plan_updates(session, base, current, row)
            if not operations:
                unchanged += 1
            # Merge into a copy: a failure must not leave half of a row's plan.
            candidate = deepcopy(updates)
            for target_id, before, changes, which in operations:
                incoming = EventUpdate(target_id, before, changes, which, [number])
                if target_id in candidate:
                    merge_update(candidate[target_id], incoming)
                else:
                    candidate[target_id] = incoming
            updates = candidate
        except (ValueError, KeyError, TypeError, AttributeError, RuntimeError, requests.RequestException) as exc:
            errors.append(f'Row {number}, event {event_id}: {exc}')
    if errors:
        raise ValueError('Import validation failed; no events were updated:\n' + '\n'.join(errors))
    # Parent schedule replacement must precede direct child-time edits. Otherwise
    # the parent's unchanged sibling times could overwrite a child's new time.
    return sorted(updates.values(), key=lambda update: 'sections' not in update.changes), unchanged


def display_update(update):
    """Print the exact targets, old/new values, and section-wide effects."""
    print(f'Rows {sorted(set(update.rows))} -> event {update.event_id}, scope={update.which}:')
    if update.before.get('child_events_count') and SHARED_SECTION_FIELDS & update.changes.keys():
        print('  Shared parent fields affect ALL sibling sections.')
    if 'sections' in update.changes:
        print('  Section replacement can change child event IDs; re-export after applying.')
    for key, value in sorted(update.changes.items()):
        old = update.before.get('child_events') if key == 'sections' else update.before.get(key)
        print(f'  {key}: {json.dumps(old)} -> {json.dumps(value)}')


def configure_parser(parser):
    parser.add_argument('spreadsheet', nargs='?', type=Path)
    parser.add_argument('--sheet', help='Excel worksheet name; otherwise use active worksheet')
    parser.add_argument('--url', default=os.getenv('CANVAS_URL'))
    parser.add_argument('--apply', action='store_true', help='Write the previewed changes to Canvas')


def parse_arguments():
    parser = argparse.ArgumentParser(description="Update existing Canvas calendar events")
    configure_parser(parser)
    return parser.parse_args()


def main(args=None):
    """Read -> validate/plan -> preview -> optionally apply; never create events."""
    if args is None:
        args = parse_arguments()
    args.spreadsheet = resolve_input_path(args.spreadsheet, DEFAULT_CSV)
    rows = read_rows(args.spreadsheet, args.sheet)
    base, session = connect(args.url)
    with session:
        print(f'Building preview for {len(rows)} rows. Fetching current Canvas events...', flush=True)
        updates, unchanged = build_import_plan(session, base, rows)
        for update in updates:
            display_update(update)
        if not args.apply:
            print(f'Preview only: {len(updates)} event requests; {unchanged} unchanged rows.')
            print('Run again with --apply to save these changes.')
            return 0
        completed = 0
        for update in updates:
            try:
                url = f'{base}/api/v1/calendar_events/{update.event_id}'
                canvas_request(session, 'PUT', url, data=build_update_payload(update.changes, update.which))
                completed += 1
                print(f'Updated event {update.event_id}')
            except (ValueError, RuntimeError, requests.RequestException) as exc:
                # Canvas offers no transaction across these requests. Stop on the
                # first failure and report any writes that already succeeded.
                print(f'FAILED event {update.event_id}: {exc}')
                print(f'{completed} requests succeeded. Preview again before retrying.')
                return 1
        print(f'Updated {completed} event requests; {unchanged} unchanged rows.')
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except (ValueError, OSError, KeyError, ImportError) as exc:
        raise SystemExit(str(exc))
